"""MDA v2: per-day/per-developer grain, real tokens from OpenTelemetry, estimate flow, date-ranged
AI Efficiency report and workbook, and the effort-accounting rules (idle gaps, midnight/weekends,
unassigned turns, forced continuation, interrupted turns).

Run: python -m unittest discover -s tests/mda -v   (stdlib only; the workbook test needs openpyxl)
"""
import datetime as dt
import json
import os
import subprocess
import sys
import unittest
from pathlib import Path

from test_mda import Base, REPO

UTC = dt.timezone.utc


def T(y, mo, d, h=0, mi=0):
    return dt.datetime(y, mo, d, h, mi, tzinfo=UTC)


def cfg_utc():
    from mda.core import load_config
    cfg = load_config()
    cfg["effort"]["timezone"] = "UTC"
    return cfg


# ----------------------------------------------------------------------------- effort rules

class TestEffortRules(Base):
    def test_weekend_gap_does_not_create_worked_days(self):
        """Fri 17:00 -> Mon 09:00 approval: idle, excluded; no Sat/Sun/Mon 'days' appear."""
        from mda import effort as E
        cfg = cfg_utc()
        E.add_entry(self.story, session_id="s", stage="build", kind="approval",
                    start=T(2026, 10, 2, 17), end=T(2026, 10, 5, 9), cfg=cfg)
        row = list(E.rollup(self.story, cfg).values())[0]
        self.assertEqual((row["approval_seconds"], row["active_days"], row["idle_excluded"]), (0, 0, 1))
        self.assertEqual(E.effort_rows(self.story, cfg)[0]["idle_excluded"], 1)

    def test_midnight_crossing_counts_both_days_actually_worked(self):
        from mda import effort as E
        cfg = cfg_utc()
        E.add_entry(self.story, session_id="s", stage="build", kind="ai",
                    start=T(2026, 10, 2, 23), end=T(2026, 10, 3, 1), cfg=cfg)
        rows = E.effort_rows(self.story, cfg)
        self.assertEqual([r["work_date"] for r in rows], ["2026-10-02", "2026-10-03"])
        self.assertAlmostEqual(sum(r["ai_seconds"] for r in rows), 7200, delta=1)

    def test_capped_turn_counts_only_its_capped_span(self):
        """A stuck turn from Fri 22:00 to Monday is capped at 2h from its start -> only Friday."""
        from mda import effort as E
        cfg = cfg_utc()
        E.add_entry(self.story, session_id="s", stage="build", kind="ai",
                    start=T(2026, 10, 2, 22), end=T(2026, 10, 5, 9), cfg=cfg)
        self.assertEqual([r["work_date"] for r in E.effort_rows(self.story, cfg)], ["2026-10-02"])

    def test_turn_outside_any_stage_is_unassigned_not_dev(self):
        from mda import effort as E
        e = E.add_entry(self.story, session_id="s", stage=None, kind="ai", start=0, end=1800)
        E.add_entry(self.story, session_id="s", stage="build", kind="ai", start=0, end=600)
        self.assertEqual((e["stage"], e["phase"]), ("unassigned", "unassigned"))
        p = E.story_totals(self.story)["active_seconds_by_phase"]
        self.assertEqual((p["dev"], p["unassigned"]), (600, 1800))

    def test_manual_phases_are_their_own_phase(self):
        from mda import effort as E
        E.add_entry(self.story, session_id="s", stage="transition_to_live", kind="manual", start=0, end=3600)
        E.add_entry(self.story, session_id="s", stage="support", kind="manual", start=0, end=1800)
        p = E.story_totals(self.story)["active_seconds_by_phase"]
        self.assertEqual((p["transition_to_live"], p["support"], p["dev"]), (3600, 1800, 0))

    def test_usage_entries_carry_tokens_but_no_time(self):
        from mda import effort as E
        E.add_entry(self.story, session_id="s", stage="build", kind="ai", start=0, end=600, est_tokens=5)
        E.add_entry(self.story, session_id="s", stage="build", kind="usage", start=0, end=500, agent="a",
                    tokens_in=900, tokens_out=100, token_source="otel")
        t = E.story_totals(self.story)
        self.assertEqual((t["active_seconds_total"], t["tokens_in"], t["tokens_out"], t["tokens_source"]),
                         (600, 900, 100, "otel"))
        self.assertEqual(t["tokens_est"], 5)  # estimate kept separate, never summed with real tokens

    def test_rows_split_by_developer(self):
        from mda import effort as E
        cfg = cfg_utc()
        E.add_entry(self.story, session_id="s", stage="build", kind="ai", start=T(2026, 10, 2, 9),
                    end=T(2026, 10, 2, 10), cfg=cfg, developer="dev1")
        E.add_entry(self.story, session_id="s", stage="build", kind="ai", start=T(2026, 10, 2, 9),
                    end=T(2026, 10, 2, 11), cfg=cfg, developer="dev2")
        self.assertEqual(len(E.effort_rows(self.story, cfg)), 2)
        mine = E.effort_rows(self.story, cfg, developer="dev2")
        self.assertEqual((len(mine), mine[0]["ai_seconds"]), (1, 7200))


# ----------------------------------------------------------------------------- hooks

class TestHookAccounting(Base):
    def run_hook(self, event, payload):
        from mda import hooks
        payload.setdefault("sessionId", "sess1")
        os.environ["MDA_DEBUG"] = "1"
        return hooks.dispatch(event, payload)

    def bind(self, stage=None):
        S = self.S
        if stage:
            S.start_stage(self.story, stage)
        S.save_story(self.story)
        sess = S.load_session("sess1")
        sess["story_id"] = "1234"
        if stage:
            sess.update({"active_stage": stage, "stage_run_open": True})
        S.save_session(sess)

    def ai_seconds(self):
        s, _ = self.S.load_story("1234")
        return [(x["stage"], x["seconds"], x.get("note")) for x in s["effort"] if x["kind"] == "ai"]

    def test_forced_continuation_time_is_not_lost(self):
        self.bind("requirements")
        t = 1_800_000_000_000
        self.run_hook("userPromptSubmitted", {"timestamp": t, "prompt": "go"})
        self.assertEqual(self.run_hook("agentStop", {"timestamp": t + 30_000}).get("decision"), "block")
        self.assertEqual(self.run_hook("agentStop", {"timestamp": t + 90_000}).get("decision"), "block")
        self.run_hook("agentStop", {"timestamp": t + 100_000})  # forced-continuation cap reached
        self.assertEqual(sum(x[1] for x in self.ai_seconds()), 100)  # 30 + 60 + 10

    def test_interrupted_turn_is_flagged_not_silently_dropped(self):
        self.bind("requirements")
        t = 1_800_000_000_000
        self.run_hook("userPromptSubmitted", {"timestamp": t, "prompt": "first"})
        self.run_hook("userPromptSubmitted", {"timestamp": t + 50_000, "prompt": "second"})  # no agentStop between
        s, _ = self.S.load_story("1234")
        self.assertIn("interrupted_turn", [a["kind"] for a in s["anomalies"]])
        self.assertIn(("requirements", 0, "interrupted_turn"), self.ai_seconds())

    def test_turn_outside_any_stage_lands_in_unassigned(self):
        self.bind()
        t = 1_800_000_000_000
        self.run_hook("userPromptSubmitted", {"timestamp": t, "prompt": "what does this do?"})
        self.run_hook("agentStop", {"timestamp": t + 20_000})
        self.assertEqual(self.ai_seconds(), [("unassigned", 20, None)])

    def test_turn_before_story_is_bound_is_not_flagged_interrupted(self):
        S = self.S
        S.save_story(self.story)
        t = 1_800_000_000_000
        self.run_hook("userPromptSubmitted", {"timestamp": t, "prompt": "hello"})   # no story bound yet
        self.run_hook("agentStop", {"timestamp": t + 5_000})
        sess = S.load_session("sess1")
        sess["story_id"] = "1234"
        S.save_session(sess)
        self.run_hook("userPromptSubmitted", {"timestamp": t + 60_000, "prompt": "now bound"})
        s, _ = S.load_story("1234")
        self.assertNotIn("interrupted_turn", [x["kind"] for x in s["anomalies"]])

    def test_activity_command_books_to_its_stage(self):
        self.bind()
        sess = self.S.load_session("sess1")
        sess["activity"] = {"name": "research", "stage": "discovery"}
        self.S.save_session(sess)
        t = 1_800_000_000_000
        self.run_hook("userPromptSubmitted", {"timestamp": t, "prompt": "/research ocr options"})
        self.run_hook("agentStop", {"timestamp": t + 20_000})
        self.assertEqual(self.ai_seconds(), [("discovery", 20, None)])


# ----------------------------------------------------------------------------- OpenTelemetry

def _attr(k, v):
    return {"key": k, "value": ({"intValue": str(v)} if isinstance(v, int) else {"stringValue": v})}


def _otlp_span(span, parent, op, start, end, agent=None, model=None, tin=None, tout=None, conv="conv-1", trace="t1"):
    attrs = [_attr("gen_ai.operation.name", op), _attr("gen_ai.conversation.id", conv)]
    if agent:
        attrs.append(_attr("gen_ai.agent.name", agent))
    if model:
        attrs.append(_attr("gen_ai.request.model", model))
    if tin is not None:
        attrs.append(_attr("gen_ai.usage.input_tokens", tin))
    if tout is not None:
        attrs.append(_attr("gen_ai.usage.output_tokens", tout))
    return {"traceId": trace, "spanId": span, "parentSpanId": parent or "", "name": f"{op} x",
            "startTimeUnixNano": str(int(start.timestamp() * 1e9)), "endTimeUnixNano": str(int(end.timestamp() * 1e9)),
            "attributes": attrs}


def otlp_line(*spans, service="copilot-chat"):
    return json.dumps({"resourceSpans": [{"resource": {"attributes": [_attr("service.name", service)]},
                                          "scopeSpans": [{"spans": list(spans)}]}]})


class TestOtel(Base):
    def write(self, *lines):
        p = self.tmp / ".mda" / "otel" / "copilot.jsonl"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return p

    def turn_lines(self, t0=T(2026, 10, 2, 10), pre=""):
        m = dt.timedelta
        k = lambda x: pre + x  # noqa: E731  (unique span/trace ids per turn)
        return [
            otlp_line(_otlp_span(k("root"), None, "invoke_agent", t0, t0 + m(minutes=5), agent="build-engineer",
                                 tin=99999, tout=99999, trace=k("t"))),  # parent totals must NOT be added on top of children
            otlp_line(_otlp_span(k("c1"), k("root"), "chat", t0, t0 + m(seconds=30), model="gpt-x", tin=1000, tout=100, trace=k("t")),
                      _otlp_span(k("c2"), k("root"), "chat", t0 + m(seconds=40), t0 + m(seconds=90), model="gpt-x",
                                 tin=2000, tout=300, trace=k("t"))),
            otlp_line(_otlp_span(k("sub"), k("root"), "invoke_agent", t0 + m(seconds=100), t0 + m(seconds=200),
                                 agent="Explore", trace=k("t")),
                      _otlp_span(k("c3"), k("sub"), "chat", t0 + m(seconds=110), t0 + m(seconds=190), model="haiku",
                                 tin=500, tout=50, trace=k("t"))),
        ]

    def test_chat_spans_roll_up_per_agent_without_double_counting(self):
        from mda import otel
        spans, stats = otel.read_spans(self.write(*self.turn_lines()))
        agg = otel.aggregate(spans)
        by = {g["agent"]: g for g in agg["usage"]}
        self.assertEqual((by["build-engineer"]["tin"], by["build-engineer"]["tout"], by["build-engineer"]["calls"]),
                         (3000, 400, 2))
        self.assertEqual((by["Explore"]["tin"], by["Explore"]["tout"]), (500, 50))  # sub-agent attributed to itself
        self.assertEqual(len(agg["turns"]), 1)  # only the root run is a turn

    def test_import_attributes_to_stage_window_else_unassigned_and_is_idempotent(self):
        from mda import effort as E, otel
        cfg = cfg_utc()
        s = self.story
        s["stage_events"] = [{"stage": "build", "event": "start", "at": "2026-10-02T09:00:00.000000Z"},
                             {"stage": "build", "event": "approve", "at": "2026-10-02T12:00:00.000000Z"}]
        p = self.write(*self.turn_lines(T(2026, 10, 2, 10)), *self.turn_lines(T(2026, 10, 2, 15), "late-")[:2])
        r = otel.import_into_story(s, cfg, p)
        self.assertTrue(r["ok"], r)
        t = E.rollup(s, cfg)
        keyed = {k[1]: v for k, v in t.items()}
        self.assertEqual((keyed["build"]["tokens_in"], keyed["build"]["tokens_out"]), (3500, 450))
        self.assertEqual(keyed["unassigned"]["tokens_in"], 3000)  # 15:00 turn is outside the stage window
        # same file again: nothing changes; forced re-import: still no duplicates
        self.assertEqual(otel.import_into_story(s, cfg, p).get("skipped"), "unchanged")
        n = len(s["effort"])
        self.assertEqual(otel.import_into_story(s, cfg, p, force=True)["usage_entries_changed"], 0)
        self.assertEqual(len(s["effort"]), n)

    def test_late_arriving_span_updates_the_same_entry(self):
        from mda import effort as E, otel
        cfg = cfg_utc()
        s = self.story
        lines = self.turn_lines()
        p = self.write(*lines[:2])
        otel.import_into_story(s, cfg, p)
        before = [x for x in s["effort"] if x["kind"] == "usage"]
        p = self.write(*lines)  # the sub-agent span was flushed later
        otel.import_into_story(s, cfg, p)
        after = [x for x in s["effort"] if x["kind"] == "usage"]
        self.assertEqual((len(before), len(after)), (1, 2))
        self.assertEqual(sum(x["tokens_in"] for x in after), 3500)

    def test_console_style_and_flat_span_shapes_are_read(self):
        from mda import otel
        t0 = T(2026, 10, 2, 10)
        flat = {"traceId": "t9", "id": "s9", "parentSpanContext": {"spanId": "r9"}, "name": "chat gpt-4o",
                "timestamp": int(t0.timestamp() * 1e6), "duration": 2_500_000,
                "attributes": {"gen_ai.operation.name": "chat", "gen_ai.usage.input_tokens": 7,
                               "gen_ai.usage.output_tokens": 3, "gen_ai.request.model": "gpt-4o"}}
        root = {"traceId": "t9", "id": "r9", "name": "invoke_agent copilot",
                "startTime": [int(t0.timestamp()), 0], "endTime": [int(t0.timestamp()) + 5, 0],
                "attributes": {"gen_ai.operation.name": "invoke_agent", "gen_ai.agent.name": "copilot"}}
        p = self.write(json.dumps(flat), json.dumps({"type": "span", "data": root}), "not json at all",
                       json.dumps({"metric": "ignored"}))
        spans, stats = otel.read_spans(p)
        self.assertEqual((len(spans), stats["unparseable_lines"]), (2, 1))
        g = otel.aggregate(spans)["usage"][0]
        self.assertEqual((g["agent"], g["tin"], g["tout"], g["model"]), ("copilot", 7, 3, "gpt-4o"))
        self.assertAlmostEqual((spans[0]["end"] - spans[0]["start"]).total_seconds(), 2.5)

    def test_turn_time_import_is_opt_in_and_never_double_counts(self):
        from mda import effort as E, otel
        cfg = cfg_utc()
        cfg["otel"]["import_turn_time"] = True
        s = self.story
        p = self.write(*self.turn_lines(T(2026, 10, 2, 10)))
        # a hook already timed that turn -> the span must not add a second AI entry
        E.add_entry(s, session_id="h", stage="build", kind="ai", start=T(2026, 10, 2, 10),
                    end=T(2026, 10, 2, 10, 5), cfg=cfg)
        otel.import_into_story(s, cfg, p)
        self.assertEqual(len([x for x in s["effort"] if x["kind"] == "ai"]), 1)
        # no hook entry -> the span supplies the turn time (VS Code Chat without turn start/end)
        s["effort"] = []
        s["otel"] = {}
        otel.import_into_story(s, cfg, p)
        ai = [x for x in s["effort"] if x["kind"] == "ai"]
        self.assertEqual((len(ai), ai[0]["seconds"]), (1, 300))

    def test_check_reports_what_a_real_file_contains(self):
        from mda import otel
        res = otel.check(self.write(*self.turn_lines()))
        self.assertTrue(res["ok"])
        self.assertEqual((res["chat_spans"], res["total_tokens_in"], res["total_tokens_out"]), (3, 3500, 450))
        self.assertFalse(otel.check(self.tmp / "missing.jsonl")["ok"])


# ----------------------------------------------------------------------------- estimate

class TestEstimate(Base):
    FOUR = {"discovery": {"pd": 4, "reason": "14 screens across 3 modules to trace"},
            "design": {"pd": 3, "reason": "new API layer"},
            "build": {"pd": 12, "reason": "screens + 4 endpoints"},
            "test": {"pd": 5, "reason": "20 scenarios"}}

    def test_llm_estimate_per_phase_with_reason_and_labels(self):
        from mda import estimate as EST
        est = EST.set_many(None, self.FOUR, "llm", counts={"screen_simple": 14})
        self.assertEqual((est["phases"]["build"], est["total_pd"], est["method"]), (12.0, 24.0, "llm"))
        self.assertIsNone(est["phases"]["transition_to_live"])          # NEEDS INPUT, never a guess
        self.assertEqual(est["entries"]["design"]["reason"], "new API layer")
        self.assertTrue(any("LLM-estimated" in x for x in est["labels"]))
        self.assertTrue(any("not yet confirmed" in x for x in est["labels"]))

    def test_reason_is_mandatory_and_bad_batch_changes_nothing(self):
        from mda import estimate as EST
        est = EST.set_many(None, {"build": {"pd": 5, "reason": "ok"}})
        with self.assertRaises(ValueError):
            EST.set_phase(est, "build", 8, "")
        bad = dict(self.FOUR, test={"pd": 1, "reason": ""})
        with self.assertRaises(ValueError):
            EST.set_many(est, bad)
        self.assertEqual(est["phases"]["build"], 5.0)                    # untouched by the rejected batch
        for args in (("nope", 1, "r"), ("build", -1, "r"), ("build", "x", "r")):
            with self.assertRaises(ValueError):
                EST.set_phase(est, *args)

    def test_developer_change_keeps_previous_value_and_resets_confirmation(self):
        from mda import estimate as EST
        est = EST.confirm(EST.set_many(None, self.FOUR), "dev1")
        self.assertTrue(est["confirmed"])
        est = EST.set_phase(est, "build", 8, "shared components already exist", "developer")
        self.assertEqual((est["phases"]["build"], est["entries"]["build"]["previous_pd"]), (8.0, 12.0))
        self.assertEqual((est["confirmed"], est["method"]), (False, "llm+developer"))

    def test_developer_supplied_transition_and_support(self):
        from mda import estimate as EST
        est = EST.set_phase(EST.set_many(None, self.FOUR), "transition_to_live", 3, "runbook + smoke test", "developer")
        self.assertEqual(est["phases"]["transition_to_live"], 3.0)
        self.assertEqual(est["total_pd"], 27.0)


# ----------------------------------------------------------------------------- AI efficiency report

def est_record(story_id, phases, tier="High", app="D2R2", stack=".Net 10, AWS", llm=True, confirmed=True):
    src = "llm" if llm else "developer"
    return {"story_id": story_id, "tier": tier, "application": app, "tech_stack": stack,
            "estimate": {"phases": phases, "confirmed": confirmed,
                         "entries": {k: {"pd": v, "source": src, "reason": "t"} for k, v in phases.items() if v is not None}}}


def row(story, day, stage, ai=0.0, appr=0.0, man=0.0, dev="d1", **kw):
    return {"story_id": story, "developer": dev, "work_date": day, "stage": stage, "cycle": 1,
            "ai_hours": ai, "approval_hours": appr, "manual_hours": man, "tokens_in": 0, "tokens_out": 0,
            "tokens_est": 0, "token_source": "none", **kw}


class TestAiEfficiencyReport(Base):
    def report(self, rows, ests, **kw):
        from mda import benefits
        return benefits.build_report("Proj", rows, ests, cfg=cfg_utc(), **kw)

    def by_phase(self, res):
        return {r["phase"]: r for r in res["rows"]}

    def test_like_for_like_phases_only_and_exclusions_listed(self):
        rows = [row("1", "2026-09-01", "discovery", ai=7.5), row("1", "2026-09-02", "build", ai=7.5)]
        res = self.report(rows, {"1": est_record("1", {"build": 10.0})})
        p = self.by_phase(res)
        self.assertIsNone(p["discovery"]["traditional_pd"])          # no traditional figure -> NEEDS INPUT
        self.assertEqual((p["build"]["traditional_pd"], p["build"]["ai_pd"]), (10.0, 1.0))
        self.assertAlmostEqual(res["benefits"]["effort_saving_pct"], 0.9)   # 1 - 1/10, discovery not mixed in
        self.assertEqual(res["excluded"], [{"story_id": "1", "phase": "discovery", "hours": 7.5}])
        self.assertIsNone(res["benefits"]["cost_saving_pct"])

    def test_date_range_filters_by_work_date_and_prorates_traditional(self):
        rows = [row("1", "2026-09-01", "build", ai=7.5), row("1", "2026-09-10", "build", ai=7.5)]
        ests = {"1": est_record("1", {"build": 10.0})}
        full = self.by_phase(self.report(rows, ests))["build"]
        self.assertEqual((full["traditional_pd"], full["ai_pd"], full["ai_days"]), (10.0, 2.0, 2))
        part = self.by_phase(self.report(rows, ests, date_from="2026-09-01", date_to="2026-09-05"))["build"]
        self.assertEqual((part["traditional_pd"], part["ai_pd"], part["ai_days"]), (5.0, 1.0, 1))
        none = self.report(rows, ests, date_from="2026-10-01", date_to="2026-10-31")
        self.assertEqual((none["story_count"], self.by_phase(none)["build"]["ai_pd"]), (0, None))

    def test_range_boundaries_are_inclusive_and_open_ends_work(self):
        rows = [row("1", d, "build", ai=7.5) for d in ("2026-09-01", "2026-09-05", "2026-09-06")]
        ests = {"1": est_record("1", {"build": 3.0})}
        self.assertEqual(self.by_phase(self.report(rows, ests, date_from="2026-09-05", date_to="2026-09-05"))["build"]["ai_days"], 1)
        self.assertEqual(self.by_phase(self.report(rows, ests, date_to="2026-09-05"))["build"]["ai_days"], 2)
        self.assertEqual(self.by_phase(self.report(rows, ests, date_from="2026-09-05"))["build"]["ai_days"], 2)

    def test_reversed_range_is_rejected(self):
        with self.assertRaises(ValueError):
            self.report([], {}, date_from="2026-09-05", date_to="2026-09-01")

    def test_ai_days_are_distinct_dates_across_developers(self):
        rows = [row("1", "2026-09-01", "build", ai=3.0, dev="a"), row("1", "2026-09-01", "build", ai=4.5, dev="b")]
        b = self.by_phase(self.report(rows, {"1": est_record("1", {"build": 5.0})}))["build"]
        self.assertEqual((b["ai_days"], b["ai_pd"]), (1, 1.0))   # 7.5 h of effort, one calendar day

    def test_stages_map_to_template_phases_and_manual_phases_report(self):
        rows = [row("1", "2026-09-01", "requirements", ai=1.5), row("1", "2026-09-01", "discovery", ai=6.0),
                row("1", "2026-09-02", "plan", ai=7.5), row("1", "2026-09-03", "quality", ai=3.75),
                row("1", "2026-09-04", "transition_to_live", man=7.5), row("1", "2026-09-05", "support", man=3.75),
                row("1", "2026-09-05", "unassigned", ai=2.0)]
        est = {"discovery": 4, "design": 2, "build": 1, "test": 1, "transition_to_live": 2, "support": 1}
        res = self.report(rows, {"1": est_record("1", dict(est, build=None))})
        p = self.by_phase(res)
        self.assertEqual(p["discovery"]["ai_pd"], 1.0)                 # requirements + discovery
        self.assertEqual((p["design"]["ai_pd"], p["test"]["ai_pd"]), (1.0, 0.5))
        self.assertEqual((p["transition_to_live"]["ai_pd"], p["support"]["ai_pd"]), (1.0, 0.5))
        self.assertTrue(any("outside any stage" in f for f in res["flags"]))

    def test_automated_completion_is_ai_share_of_active_time(self):
        rows = [row("1", "2026-09-01", "build", ai=3.0, appr=1.0)]
        res = self.report(rows, {"1": est_record("1", {"build": 2.0})})
        self.assertAlmostEqual(res["benefits"]["automated_completion_pct"], 0.75)

    def test_header_application_complexity_and_stack(self):
        rows = [row("1", "2026-09-01", "build", ai=1), row("2", "2026-09-01", "build", ai=1)]
        ests = {"1": est_record("1", {"build": 1}, tier="High"), "2": est_record("2", {"build": 1}, tier="Low", stack="SQL Server")}
        h = self.report(rows, ests)["header"]
        self.assertEqual(h["application"], "D2R2")
        self.assertTrue(h["overall_complexity"].startswith("High (mixed"))
        self.assertEqual(h["technology_stack"], ".Net 10, AWS, SQL Server")
        h2 = self.report(rows, ests, application="Other", tech_stack="X")["header"]
        self.assertEqual((h2["application"], h2["technology_stack"]), ("Other", "X"))

    def test_flags_for_llm_estimate_unconfirmed_and_missing_tokens(self):
        rows = [row("1", "2026-09-01", "build", ai=1)]
        res = self.report(rows, {"1": est_record("1", {"build": 1}, llm=True, confirmed=False)})
        text = " ".join(res["flags"])
        self.assertIn("LLM-estimated", text)
        self.assertIn("not been confirmed", text)
        self.assertIn("No real token data", text)

    def test_tokens_by_stage_and_agent_real_vs_estimate(self):
        rows = [row("1", "2026-09-01", "build", ai=1, tokens_in=900, tokens_out=100, tokens_est=40,
                    token_source="otel", agents={"build-engineer": {"in": 900, "out": 100}})]
        t = self.report(rows, {"1": est_record("1", {"build": 1})})["tokens"]
        self.assertEqual((t["total_in"], t["total_out"], t["total_est"]), (900, 100, 40))
        self.assertEqual(t["by_agent"], [{"agent": "build-engineer", "tokens_in": 900, "tokens_out": 100}])

    def test_empty_input_does_not_crash_and_renders(self):
        from mda import benefits
        res = self.report([], {})
        self.assertIsNone(res["benefits"]["effort_saving_pct"])
        self.assertIn("NEEDS INPUT", benefits.markdown(res))

    def test_from_records_reads_ado_payloads_for_one_project(self):
        from mda import benefits
        from mda.core import load_config
        fields = load_config()["ado"]["fields"]
        mk = lambda title, payload: {"System.Title": title, fields["payload"]: json.dumps(payload)}  # noqa: E731
        recs = [mk("MDA-1-d1-2026-09-01-c1-build", {"record_type": "effort", "project": "Proj", "story_id": 1,
                                                   "work_date": "2026-09-01", "stage": "build", "ai_hours": 2}),
                mk("MDA-1-estimate", {"record_type": "estimate", "project": "Proj", "story_id": 1,
                                      "estimate": {"phases": {"build": 4}}}),
                mk("MDA-2-d1-2026-09-01-c1-build", {"record_type": "effort", "project": "Other", "story_id": 2,
                                                   "work_date": "2026-09-01", "stage": "build", "ai_hours": 9}),
                mk("MDA-1-state-d1", {"record_type": "state", "project": "Proj", "story_id": 1}),
                {"System.Title": "MDA-1-d1-2026-09-02-c1-build", "System.Description": "<pre>" + json.dumps(
                    {"project": "Proj", "story_id": 1, "work_date": "2026-09-02", "stage": "build", "ai_hours": 1}) + "</pre>"}]
        rows, ests = benefits.from_records(recs, fields, "Proj")
        self.assertEqual([r["ai_hours"] for r in rows], [2, 1])            # other project + state record ignored
        self.assertIn("1", ests)

    def test_workbook_matches_template_and_has_live_formulas(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        from mda import xlsx
        rows = [row("1", "2026-09-01", "build", ai=7.5, tokens_in=10, tokens_out=5, token_source="otel")]
        res = self.report(rows, {"1": est_record("1", {"build": 10.0})}, date_from="2026-09-01", date_to="2026-09-30")
        path = xlsx.write_workbook(res, self.tmp / "out.xlsx")
        ws = openpyxl.load_workbook(path)["AI Efficiency Report"]
        self.assertEqual(ws["A1"].value, "AI Productivity (AI Efficiency Report)")
        self.assertEqual([ws.cell(row=r, column=1).value for r in (4, 5, 6)],
                         ["Application", "Overall complexity", "Technology stack"])
        self.assertEqual([ws.cell(row=10, column=c).value for c in (1, 3, 5, 7)],
                         ["Effort saving", "Time saving", "Cost saving", "Automated completion"])
        self.assertEqual([ws.cell(row=14, column=c).value for c in range(1, 10)][:5],
                         ["Phase", "Traditional effort (PD)", "AI effort (PD)", "Traditional days*", "AI days*"])
        labels = [ws.cell(row=r, column=1).value for r in range(15, 21)]
        self.assertEqual(labels, ["Discovery", "Design", "Build", "Test", "Transition to Live", "Support"])
        total = ws["A21"].value
        self.assertEqual(total, "Total / average")
        self.assertEqual(ws["B21"].value, "=SUM(B15:B20)")
        self.assertIn("1-C21/B21", ws["A11"].value)
        self.assertEqual((ws["B17"].value, ws["C17"].value), (10.0, 1.0))
        self.assertIsNone(ws["F17"].value)   # cost left blank on purpose
        self.assertIn("Details", openpyxl.load_workbook(path).sheetnames)


# ----------------------------------------------------------------------------- ADO client

class TestAdoClient(Base):
    def client(self, ft):
        os.environ["MDA_ADO_ORG_URL"] = "https://dev.azure.com/org"
        from mda.ado import AdoClient
        return AdoClient(project="Proj", pat="x", transport=ft)

    def test_selftest_proves_type_and_fields_then_cleans_up(self):
        from test_mda import FakeTransport
        ft = FakeTransport()
        res = self.client(ft).selftest()
        self.assertEqual((res["ok"], res["structured_fields"], res["cleaned_up"]), (True, True, True))
        self.assertEqual(ft.items, {})                           # the test record is deleted again

    def test_selftest_reports_when_custom_fields_are_missing(self):
        from test_mda import FakeTransport
        ft = FakeTransport()
        ft.reject_custom = True
        res = self.client(ft).selftest()
        self.assertEqual((res["ok"], res["structured_fields"]), (True, False))
        self.assertIn("Description", res["advice"])

    def test_record_is_still_written_if_the_parent_link_is_refused(self):
        from test_mda import FakeTransport
        ft = FakeTransport()
        ft.reject_links = True
        made = self.client(ft).upsert_record("MDA-1-x", {"project": "Proj"}, "{}", story_ref=1234)
        self.assertTrue(made["created"])
        self.assertFalse(made["linked"])
        self.assertEqual(len(ft.items), 1)

    def test_link_is_added_when_allowed(self):
        from test_mda import FakeTransport
        made = self.client(FakeTransport()).upsert_record("MDA-1-x", {"project": "Proj"}, "{}", story_ref=1234)
        self.assertTrue(made["linked"])


# ----------------------------------------------------------------------------- CLI end to end

class TestCliEndToEnd(Base):
    def cli(self, *args, ok=True):
        env = {**os.environ, "MDA_ROOT": str(self.tmp), "MDA_SESSION_ID": "cli-sess", "PYTHONIOENCODING": "utf-8"}
        for k in ("MDA_ADO_PAT", "MDA_ADO_ORG_URL"):
            env.pop(k, None)
        p = subprocess.run([sys.executable, str(REPO / ".github/scripts/mda.py"), *args], capture_output=True,
                           text=True, env=env, cwd=str(self.tmp), encoding="utf-8")
        try:
            out = json.loads(p.stdout)
        except ValueError:
            self.fail(f"not JSON ({p.returncode}): {p.stdout!r} {p.stderr!r}")
        if ok:
            self.assertEqual(p.returncode, 0, out)
        return out, p.returncode

    def test_estimate_manual_activity_flow(self):
        (self.tmp / ".git").mkdir()
        self.cli("setup", "set", "--project", "Proj", "--story-id", "77", "--migration", "no",
                 "--application", "D2R2", "--tech-stack", ".Net 10, AWS")
        s, _ = self.S.load_story("77")
        s["setup"]["account"] = {"confirmed_login": "dev1"}
        self.S.save_story(s)
        # discovery counts -> plan score (tier only) -> the LLM then records its own estimate
        self.cli("complexity", "counts", "--json", json.dumps({"screen_simple": 10, "internal_api_endpoint": 4}))
        out, _ = self.cli("complexity", "score")
        self.assertIsNone(out["estimate"])                        # no rate card: nothing is computed for you
        est_json = json.dumps({"discovery": {"pd": 2, "reason": "10 screens to trace"},
                               "build": {"pd": 9, "reason": "10 screens, 4 endpoints"}})
        out, _ = self.cli("estimate", "set", "--source", "llm", "--json", est_json)
        self.assertEqual((out["estimate"]["total_pd"], out["estimate"]["confirmed"]), (11.0, False))
        self.assertEqual(out["estimate"]["counts"]["screen_simple"], 10)   # discovery evidence kept with it
        out, rc = self.cli("estimate", "set", "--phase", "design", "--pd", "3", ok=False)
        self.assertEqual(rc, 2)                                   # a written reason is mandatory
        out, rc = self.cli("estimate", "set", "--json", "{not json", ok=False)
        self.assertEqual(rc, 2)
        out, _ = self.cli("estimate", "set", "--phase", "transition_to_live", "--pd", "3",
                          "--reason", "cutover runbook", "--source", "developer")
        out, _ = self.cli("estimate", "confirm")
        self.assertTrue(out["estimate"]["confirmed"])
        self.assertEqual(out["estimate"]["confirmed_by"], "dev1")
        # manual effort for a phase the agent does not run
        out, _ = self.cli("manual", "log", "--phase", "transition_to_live", "--minutes", "90", "--note", "cutover")
        _, rc = self.cli("manual", "log", "--phase", "build", "--minutes", "5", ok=False)
        self.assertEqual(rc, 2)
        # activity booking
        out, _ = self.cli("activity", "start", "--name", "research")
        self.assertEqual(out["booked_to_stage"], "discovery")
        self.cli("activity", "end")
        _, rc = self.cli("activity", "start", "--name", "bogus", ok=False)
        self.assertEqual(rc, 2)
        # local preview for the AI efficiency report incl. xlsx, with a date range
        out, _ = self.cli("benefits", "--project", "Proj", "--from", "2000-01-01", "--to", "2099-12-31", "--xlsx")
        t2l = next(r for r in out["rows"] if r["phase"] == "transition_to_live")
        self.assertEqual((t2l["traditional_pd"], t2l["ai_pd"]), (3.0, 0.2))     # 90 min = 0.2 PD
        self.assertTrue(any("LLM-estimated" in f for f in out["flags"]))
        self.assertTrue(Path(out["xlsx"]).exists())
        bad, rc = self.cli("benefits", "--project", "Proj", "--from", "2026-02-01", "--to", "2026-01-01", ok=False)
        self.assertEqual(rc, 2)

    def test_unknown_project_is_rejected_when_projects_are_configured(self):
        (self.tmp / ".git").mkdir()
        (self.tmp / ".mda").mkdir()
        (self.tmp / ".mda" / "config.local.json").write_text(json.dumps(
            {"ado": {"projects": [{"name": "TeamA"}, {"name": "TeamB"}]}}))
        out, rc = self.cli("setup", "set", "--project", "Nope", "--story-id", "5", ok=False)
        self.assertEqual((rc, out["reason"], out["allowed"]), (2, "unknown_ado_project", ["TeamA", "TeamB"]))
        out, rc = self.cli("setup", "set", "--project", "teamb", "--story-id", "5")
        self.assertEqual(out["project"], "teamb")

    def test_doctor_lists_every_check_and_fails_without_pat(self):
        (self.tmp / ".git").mkdir()
        out, rc = self.cli("doctor", ok=False)
        names = {c["check"] for c in out["checks"]}
        self.assertTrue({"python", "git", "hooks file", "ADO PAT", "ADO org URL", "Copilot OpenTelemetry"} <= names)
        self.assertEqual(rc, 3)                                   # no PAT in the test environment
        pat = next(c for c in out["checks"] if c["check"] == "ADO PAT")
        self.assertEqual(pat["status"], "fail")
        self.assertIn("fix", pat)

    def test_ado_selftest_cli_needs_configuration(self):
        (self.tmp / ".git").mkdir()
        out, rc = self.cli("ado", "selftest", ok=False)
        self.assertEqual((rc, out["reason"]), (3, "ado_not_configured"))

    def test_otel_check_cli_reports_missing_then_present(self):
        (self.tmp / ".git").mkdir()
        out, rc = self.cli("otel", "check", ok=False)
        self.assertEqual((rc, out["reason"]), (3, "file_missing"))
        p = self.tmp / ".mda" / "otel" / "copilot.jsonl"
        p.parent.mkdir(parents=True)
        p.write_text(otlp_line(_otlp_span("c", None, "chat", T(2026, 10, 2, 10), T(2026, 10, 2, 10, 1),
                                          model="m", tin=11, tout=2)) + "\n", encoding="utf-8")
        out, rc = self.cli("otel", "check")
        self.assertEqual((out["chat_spans"], out["total_tokens_in"]), (1, 11))


if __name__ == "__main__":
    unittest.main()
