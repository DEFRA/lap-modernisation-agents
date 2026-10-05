"""Excel workbook for the AI Productivity (AI Efficiency) report, laid out like the programme template:
title, Application Context, Benefits, Delivery Comparison (live formulas for the totals and
percentages), plus a Details sheet (range, tokens, flags, definitions) and the raw Effort Rows.

Needs `openpyxl` (installed by the pipeline). Cost cells are left blank on purpose.
"""
from __future__ import annotations

from pathlib import Path

NAVY, CYAN, CREAM, PALE = "1B365D", "00A3C4", "FFFDF0", "E8F4F8"
POUND = "£"


def write_workbook(result: dict, path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "AI Efficiency Report"

    def fill(color):
        return PatternFill("solid", start_color=color, end_color=color)

    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_bold = Font(bold=True, color="FFFFFF", name="Calibri")

    def band(row, text, size=11):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        c = ws.cell(row=row, column=1, value=text)
        c.font, c.fill = Font(bold=True, color="FFFFFF", size=size, name="Calibri"), fill(NAVY)
        c.alignment = Alignment(vertical="center")

    band(1, "AI Productivity (AI Efficiency Report)", 16)
    ws.row_dimensions[1].height = 30
    h = result["header"]
    rng = f"{result['date_from'] or 'start of data'} to {result['date_to'] or 'latest'}"

    band(3, "Application Context")
    for i, (label, val) in enumerate([("Application", h["application"]),
                                      ("Overall complexity", h["overall_complexity"]),
                                      ("Technology stack", h["technology_stack"]),
                                      ("Date range", rng)]):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        c = ws.cell(row=r, column=3, value=val if val else "NEEDS INPUT")
        c.fill, c.font = fill(CREAM), Font(color="0000FF")
        for col in (1, 2, 3, 4, 5):
            ws.cell(row=r, column=col).border = box

    band(9, "Benefits")
    heads = ["Effort saving", "Time saving", "Cost saving", "Automated completion"]
    spans = [(1, 2), (3, 4), (5, 6), (7, 9)]
    first = 15
    last = first + len(result["rows"]) - 1
    total = last + 1
    vals = [f'=IF(B{total}>0,1-C{total}/B{total},"")', f'=IF(D{total}>0,1-E{total}/D{total},"")', None,
            result["benefits"]["automated_completion_pct"]]
    for (a, b), head, val in zip(spans, heads, vals):
        for row, v in ((10, head), (11, val)):
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
            c = ws.cell(row=row, column=a, value=v)
            c.alignment = Alignment(horizontal="center", vertical="center")
            if row == 10:
                c.font, c.fill = white_bold, fill(CYAN)
            else:
                c.font, c.fill = Font(bold=True, size=18, color=NAVY), fill(PALE)
                c.number_format = "0.0%"
    ws.row_dimensions[11].height = 30

    band(13, "DELIVERY COMPARISON")
    cols = ["Phase", "Traditional effort (PD)", "AI effort (PD)", "Traditional days*", "AI days*",
            f"Traditional cost ({POUND})", f"AI cost ({POUND})", "Automation %", "Evidence link/notes"]
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=14, column=i, value=name)
        c.font, c.fill = white_bold, fill(CYAN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    ws.row_dimensions[14].height = 32
    for i, r in enumerate(result["rows"]):
        row = first + i
        cells = [r["label"], r["traditional_pd"], r["ai_pd"], r["traditional_days"], r["ai_days"], None, None,
                 f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(C{row}),B{row}>0),1-C{row}/B{row},"")', r["evidence"]]
        for col, v in enumerate(cells, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = box
            c.fill = fill(CREAM)
            if col in (2, 3, 4, 5):
                c.font, c.number_format = Font(color="0000FF"), "0.00"
            if col == 8:
                c.number_format = "0.0%"
            if col == 9:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
    ws.cell(row=total, column=1, value="Total / average")
    for col, letter in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        ws.cell(row=total, column=col, value=f"=SUM({letter}{first}:{letter}{last})").number_format = "0.00"
    ws.cell(row=total, column=8, value=f'=IF(B{total}>0,1-C{total}/B{total},"")').number_format = "0.0%"
    for col in range(1, 10):
        c = ws.cell(row=total, column=col)
        c.font, c.fill, c.border = Font(bold=True), fill(PALE), box
    ws.cell(row=total + 2, column=1,
            value="* PD = total effort in person-days. Days = distinct calendar days actually worked (AI, measured); "
                  "Traditional days = traditional PD. Phases without both figures are blank (NEEDS INPUT).")
    for note_row, f in enumerate(result["flags"], total + 3):
        ws.cell(row=note_row, column=1, value=f"! {f}").font = Font(color="C00000")
    for letter, w in zip("ABCDEFGHI", (24, 14, 12, 14, 10, 14, 12, 13, 60)):
        ws.column_dimensions[letter].width = w

    # ---- Details sheet
    d = wb.create_sheet("Details")
    pos = {"r": 1}

    def put(*vals, bold=False):
        for i, v in enumerate(vals, 1):
            d.cell(row=pos["r"], column=i, value=v).font = Font(bold=bold)
        pos["r"] += 1

    def gap():
        pos["r"] += 1

    put("Project", result["project"], bold=True)
    put("Date range", rng)
    put("Generated", result["generated_at"])
    put("Stories included", result["story_count"])
    put("Story IDs", ", ".join(result["stories"]))
    put("Model", result["model_version"])
    put("Hours per day", result["hours_per_day"])
    gap()
    put("Data-quality flags", bold=True)
    for f in result["flags"] or ["none"]:
        put(f)
    gap()
    t = result["tokens"]
    put("Tokens by stage (real = OpenTelemetry)", bold=True)
    put("Stage", "Input tokens", "Output tokens", "Prompt-size estimate (not usage)", bold=True)
    for x in t["by_stage"]:
        put(x["stage"], x["tokens_in"], x["tokens_out"], x["tokens_est"])
    put("Total", t["total_in"], t["total_out"], t["total_est"], bold=True)
    gap()
    put("Tokens by agent", bold=True)
    put("Agent", "Input tokens", "Output tokens", bold=True)
    for x in t["by_agent"]:
        put(x["agent"], x["tokens_in"], x["tokens_out"])
    gap()
    put("Token source mix (rows)", ", ".join(f"{k}: {v}" for k, v in t["source_mix"].items()) or "none")
    gap()
    put("Excluded from the comparison (no traditional figure for the phase)", bold=True)
    put("Story", "Phase", "Hours", bold=True)
    for x in result["excluded"] or []:
        put(x["story_id"], x["phase"], x["hours"])
    gap()
    put("Definitions", bold=True)
    for line in result["definitions"]:
        put(line)
    d.column_dimensions["A"].width = 44
    for letter in "BCD":
        d.column_dimensions[letter].width = 22

    # ---- raw rows
    e = wb.create_sheet("Effort Rows")
    keys = ["work_date", "story_id", "developer", "cycle", "stage", "phase", "ai_hours", "approval_hours",
            "manual_hours", "tokens_in", "tokens_out", "tokens_est", "token_source", "capped", "idle_excluded"]
    e.append(keys)
    for row in sorted(result["effort_rows"],
                      key=lambda x: (x["work_date"], str(x["story_id"]), str(x.get("developer")))):
        e.append([row.get(k) for k in keys])
    for c in e[1]:
        c.font = Font(bold=True)
    e.freeze_panes = "A2"
    wb.save(path)
    return path
